"""Real P-04 parsers and encrypted publication; no live model server required.

Optional document dependencies and explicitly prepared local OCR/PDF models
are required. Public HWP/HWPX fixtures are prepared separately, never fetched
by an ordinary pytest run. Their source URLs are in the P-04 evidence.
"""
from __future__ import annotations

import base64
import io
import json
import os
import stat
import struct
import sys
import zipfile
from dataclasses import asdict
from pathlib import Path
from unittest.mock import patch

import pytest

pytest.importorskip('docling')
from airag import api, embeddings, normalize, profiles, store
from airag import document_adapters as da
from airag import image_adapters as ia
from airag.db import connect

TOKENIZER = '97b0c614be4d77ee51c0cef4e5f07c00f9eb65b3'
SHARE = Path(sys.prefix)/'share'


def config(**kwargs):
    return {'adapter':'documents_v1',**kwargs}


def docx_bytes():
    from docx import Document
    from PIL import Image
    doc=Document();doc.add_heading('Inventory',0)
    doc.add_paragraph('Repeated source paragraph')
    doc.add_paragraph('Repeated source paragraph')
    table=doc.add_table(rows=2,cols=2)
    for cell,text in zip([table.cell(0,0),table.cell(0,1),table.cell(1,0),table.cell(1,1)],['Item','Quantity','Coffee','42']):cell.text=text
    png=io.BytesIO();Image.new('RGB',(20,20),'red').save(png,format='PNG');doc.add_picture(io.BytesIO(png.getvalue()))
    out=io.BytesIO();doc.save(out);return out.getvalue(),png.getvalue()


def xlsx_bytes():
    from openpyxl import Workbook
    wb=Workbook();a=wb.active;a.title='Coffee';a.append(['Item','Count']);a.append(['Beans',42]);a['C2']='=B2*2';a.merge_cells('A4:B4');a['A4']='Merged heading'
    b=wb.create_sheet('Tea');b.append(['Item','Count']);b.append(['Leaves',17])
    out=io.BytesIO();wb.save(out);wb.close();return out.getvalue()


def image_bytes():
    from PIL import Image,ImageDraw,ImageFont
    path=os.getenv('AIRAG_TEST_FONT','/System/Library/Fonts/Supplemental/Arial Unicode.ttf')
    if not Path(path).is_file():pytest.skip('set AIRAG_TEST_FONT to a Korean-capable TrueType font')
    image=Image.new('RGB',(1050,280),'white');d=ImageDraw.Draw(image);font=ImageFont.truetype(path,44)
    lines=['고양이는 잠을 많이 잡니다.','서울의 날씨는 맑습니다.','커피 원두를 곱게 갈아 주세요.']
    for i,line in enumerate(lines):d.text((40,25+i*75),line,font=font,fill='black')
    out=io.BytesIO();image.save(out,format='PNG');return out.getvalue(),'\n'.join(lines)


def pdf_bytes(*, mixed=False):
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader
    out=io.BytesIO();c=canvas.Canvas(out)
    c.drawString(60,780,'Coffee inventory')
    for ri,row in enumerate([['Item','Quantity'],['Coffee','42'],['Tea','17']]):
        y=740-ri*30
        for ci,text in enumerate(row):c.drawString(65+ci*180,y,text)
        c.line(60,y-8,420,y-8)
    for x in [60,240,420]:c.line(x,752,x,672)
    if mixed:
        c.showPage();data,_=image_bytes();c.drawImage(ImageReader(io.BytesIO(data)),40,500,width=525,height=140)
    c.save();return out.getvalue()


def pdf_config(**kw):
    path=Path(os.getenv('AIRAG_DOCLING_MODELS',str(SHARE/'docling-models')))
    if not path.is_dir():pytest.skip('prepare Docling layout/table models and set AIRAG_DOCLING_MODELS')
    return config(pdf_artifacts_path=str(path),**kw)


def sample(name):
    path=Path(os.getenv('AIRAG_HWP_SAMPLES',str(SHARE/'airag/samples')))/name
    if not path.is_file():pytest.skip('prepare the documented public HWP/HWPX fixtures')
    return path.read_bytes()


def hwp_config():
    path=Path(os.getenv('AIRAG_HWP_JAR',str(SHARE/'airag/hwplib-1.1.11.jar')))
    if not path.is_file():pytest.skip('set AIRAG_HWP_JAR to the pinned hwplib jar')
    return config(hwp_classpath=str(path))


def test_document_profile_captures_dependencies_and_rejects_drift():
    for bad in (True,6.0,['6']):
        with pytest.raises(ValueError):da.canonical_config(config(tesseract_psm=bad))
    cfg=da.canonical_config(config())
    assert cfg['dependency_versions']['docling-slim']=='2.126.0'
    cfg['dependency_versions']['docling-slim']='other'
    with pytest.raises(da.DocumentError,match='differ'):da.canonical_config(cfg)
    with pytest.raises(da.DocumentError):normalize.build_document(da.DOCX,docx_bytes()[0],name='doc')


def test_docx_structure_duplicate_paragraph_anchors_and_original_picture_bytes():
    content,picture=docx_bytes();nd=normalize.build_document(da.DOCX,content,name='doc',config=config())
    assert len(nd.doc.tables)==1 and 'Coffee' in nd.doc.export_to_markdown()
    duplicate=[a for a in nd.anchors.values() if a.original_text=='Repeated source paragraph']
    assert len(duplicate)==2 and duplicate[0].locator['xpath']!=duplicate[1].locator['xpath']
    from lxml import etree
    with zipfile.ZipFile(io.BytesIO(content)) as z:
        for a in duplicate:
            root=etree.fromstring(z.read(a.locator['member']))
            element=root.xpath(a.locator['xpath'],namespaces=a.locator['namespaces'])[0]
            assert ''.join(element.itertext())==a.original_text
    assert any(base64.b64decode(r['content_base64'])==picture for r in nd.artifacts['resources'])


def test_xlsx_sheet_formula_cached_value_and_merged_cell_contract():
    nd=normalize.build_document(da.XLSX,xlsx_bytes(),name='workbook',config=config())
    assert len(nd.doc.tables)>=2
    assert {a.locator.get('sheet') for a in nd.anchors.values()} >= {'Coffee','Tea'}
    sheet=nd.artifacts['sheets'][0]
    assert sheet['merged_ranges']==['A4:B4']
    formula=next(c for c in sheet['cells'] if c['cell']=='C2')
    assert formula['formula']=='=B2*2' and formula['cached_value'] is None and formula['recalculated'] is False
    assert 'Coffee!C2: =B2*2' in nd.doc.export_to_markdown()


def test_pdf_table_and_page_anchors():
    nd=normalize.build_document('application/pdf',pdf_bytes(),name='pdf',config=pdf_config())
    assert len(nd.doc.tables)==1
    assert 'Coffee' in nd.doc.export_to_markdown() and '42' in nd.doc.export_to_markdown()
    assert all(a.locator['provenance'][0]['page_no']==1 for a in nd.anchors.values())
    assert all(a.kind=='document_location' and a.granularity=='page_region' for a in nd.anchors.values())


def test_mixed_digital_scanned_pdf_and_private_ocr_temp_cleanup(monkeypatch):
    import tempfile
    created=[];real=tempfile.NamedTemporaryFile
    def temporary(*args,**kwargs):
        f=real(*args,**kwargs)
        if kwargs.get('suffix')=='.png':
            assert stat.S_IMODE(os.stat(f.name).st_mode)==0o600
            created.append(f.name)
        return f
    monkeypatch.setattr(tempfile,'NamedTemporaryFile',temporary)
    nd=normalize.build_document('application/pdf',pdf_bytes(mixed=True),name='mixed',config=pdf_config())
    text=nd.doc.export_to_markdown()
    # TSV word boundaries can separate Hangul syllables. Preserve the raw
    # recognized text, and compare content independently of word spacing.
    compact=''.join(text.split())
    assert 'Coffee' in text and '고양이' in compact and '커피' in compact
    assert {p['page_no'] for a in nd.anchors.values() for p in a.locator['provenance']}=={1,2}
    assert created and all(not Path(p).exists() for p in created)


@pytest.mark.parametrize('engine',['tesseract','easyocr'])
def test_korean_ocr_has_observed_regions_and_never_claims_original_byte_offsets(engine):
    model=SHARE/'docling-models/EasyOcr'
    if engine=='easyocr' and not model.is_dir():pytest.skip('prepare EasyOCR Korean models')
    cfg=config(image_mode='ocr',ocr_engine=engine,ocr_model_path=str(model))
    image, expected=image_bytes()
    nd=normalize.build_document('image/png',image,name='scan',config=cfg)
    recognized=''.join(nd.doc.export_to_markdown().split())
    assert '고양이' in recognized and '커피' in recognized
    assert nd.anchors and all(a.text_origin=='ocr' and a.start is None for a in nd.anchors.values())
    assert all(a.locator.get('bbox') or a.locator.get('polygon') for a in nd.anchors.values())


def test_hwp_actual_table_and_picture_extraction():
    nd=normalize.build_document('application/x-hwp',sample('표.hwp'),name='hwp',config=hwp_config())
    assert len(nd.doc.tables)==2 and 'ABC' in nd.doc.export_to_markdown()
    assert all(a.locator['path'].startswith('section[') for a in nd.anchors.values())
    image=normalize.build_document('application/x-hwp',sample('그림.hwp'),name='picture',config=hwp_config())
    from PIL import Image
    assert image.artifacts['resources']
    for r in image.artifacts['resources']:
        data=base64.b64decode(r['content_base64']);assert len(data)==r['byte_length']
        with Image.open(io.BytesIO(data)) as im:assert im.width>0


def test_hwpx_actual_table_and_picture_extraction():
    nd=normalize.build_document('application/hwp+zip',sample('SimpleTable.hwpx'),name='hwpx',config=config())
    assert len(nd.doc.tables)==1
    assert nd.doc.tables[0].data.num_rows==3 and nd.doc.tables[0].data.num_cols==3
    picture=normalize.build_document('application/hwp+zip',sample('SimplePicture.hwpx'),name='picture',config=config())
    assert picture.artifacts['resources']
    with zipfile.ZipFile(io.BytesIO(sample('SimplePicture.hwpx'))) as z:
        for r in picture.artifacts['resources']:assert base64.b64decode(r['content_base64'])==z.read(r['member'])


@pytest.mark.parametrize('media',[da.DOCX,da.XLSX,'application/pdf','application/x-hwp','application/hwp+zip','image/png'])
def test_malformed_binary_content_is_a_structured_failure(media):
    cfg=pdf_config() if media=='application/pdf' else hwp_config() if media=='application/x-hwp' else config(image_mode='ocr')
    with pytest.raises(store.ApiError):normalize.build_document(media,b'not a real document',name='bad',config=cfg)


def test_hwpx_rejects_missing_spine_coverage_and_external_entities():
    original=sample('SimpleTable.hwpx')
    for replacement in [b'<broken>',b'<!DOCTYPE x [<!ENTITY stolen SYSTEM "file:///etc/passwd">]><x>&stolen;</x>']:
        out=io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(original)) as src,zipfile.ZipFile(out,'w') as dst:
            for n in src.namelist():dst.writestr(n,replacement if n=='Contents/content.hpf' else src.read(n))
        with pytest.raises(store.ApiError):normalize.build_document('application/hwp+zip',out.getvalue(),name='bad',config=config())


def test_expansion_limits_checked_before_office_conversion():
    with pytest.raises(store.ApiError):normalize.build_document(da.DOCX,docx_bytes()[0],name='bounded',config=config(max_expanded_bytes=16))
    with pytest.raises(store.ApiError):normalize.build_document('image/png',image_bytes()[0],name='bounded',config=config(image_mode='ocr',max_image_pixels=16))


def test_document_pipeline_publication_encryption_original_roundtrip_and_failed_update(tmp_path,key_bytes):
    con=connect.init_db(tmp_path/'documents.db',key_bytes)
    content,_=docx_bytes()
    embed=embeddings.EmbeddingSpaceConfig('openai-compatible','http://localhost:9999/v1','fixture','fixture-v1',4,tokenizer_revision=TOKENIZER)
    try:
        registered=api.register_profile(con,{'embedding':asdict(embed),'normalizer':config()});assert registered['ok'],registered
        profile=registered['data']['profile_key']
        with patch('airag.embeddings.embed_documents',side_effect=lambda cfg,texts,**kwargs:[struct.pack('4f',1,0,0,0) for _ in texts]):
            added=api.add(con,'doc-add','docs','report.docx',content,media_type=da.DOCX,profile_key=profile)
        assert added['ok'] and added['data']['searchable'],added
        artifact=api.fetch_artifact(con,'docs','report.docx',generation_id=added['data']['generation_id'])['data']
        assert artifact['structured']['normalization']['artifacts']['resources']
        fetched=api.fetch_original(con,'docs','report.docx',include_bytes=True)['data']
        assert base64.b64decode(fetched['content_base64'])==content
        failed=api.update(con,'doc-fail','docs','report.docx',added['data']['revision_token'],b'broken',media_type=da.DOCX,profile_key=profile)
        assert failed['ok'] is False
        assert api.status(con,'docs','report.docx')['data']['ready_generation_id']==added['data']['generation_id']
        for p in tmp_path.iterdir():
            if p.is_file():assert b'Repeated source paragraph' not in p.read_bytes()
    finally:con.close()


def test_captured_model_digest_cannot_be_silently_replaced():
    cfg=da.canonical_config(pdf_config())
    assert cfg['model_fingerprints']
    key=next(iter(cfg['model_fingerprints']))
    cfg['model_fingerprints'][key]='0'*64
    with pytest.raises(da.DocumentError,match='artifacts differ'):da.canonical_config(cfg)


def test_sparse_huge_workbook_is_rejected_before_docling_allocation():
    from openpyxl import Workbook
    wb=Workbook();wb.active['XFD1048576']='out of bounds'
    b=io.BytesIO();wb.save(b);wb.close()
    with patch('airag.document_adapters._converter') as converter:
        with pytest.raises(store.ApiError):normalize.build_document(da.XLSX,b.getvalue(),name='large',config=config())
    converter.assert_not_called()


def test_complex_formula_is_refused_instead_of_becoming_a_python_object_repr():
    from openpyxl import Workbook
    from openpyxl.worksheet.formula import ArrayFormula
    wb=Workbook();wb.active['A1']='Numbers';wb.active['B2']=ArrayFormula(ref='B2:B3',text='=A2:A3*2')
    out=io.BytesIO();wb.save(out);wb.close()
    with pytest.raises(da.DocumentError,match='array/data-table'):
        normalize.build_document(da.XLSX,out.getvalue(),name='array',config=config())


@pytest.mark.parametrize('format,media,frames',[('JPEG','image/jpeg',1),('BMP','image/bmp',1),('TIFF','image/tiff',2)])
def test_decoded_image_formats_keep_frame_identity(format,media,frames):
    from PIL import Image
    first=Image.new('RGB',(30,20),'white');out=io.BytesIO()
    if frames==2:first.save(out,format=format,save_all=True,append_images=[Image.new('RGB',(30,20),'black')])
    else:first.save(out,format=format)
    with patch('airag.image_adapters.recognize',return_value=[{'text':'fixture','bbox':[0,0,20,10],'confidence':1}]):
        nd=normalize.build_document(media,out.getvalue(),name='frames',config=config(image_mode='ocr'))
    assert len(nd.artifacts['frames'])==frames
    assert {a.locator['frame'] for a in nd.anchors.values()}==set(range(frames))


def test_declared_binary_media_mismatch_is_rejected():
    with pytest.raises(store.ApiError):normalize.build_document(da.DOCX,xlsx_bytes(),name='mismatch',config=config())
    with pytest.raises(store.ApiError):normalize.build_document('image/jpeg',image_bytes()[0],name='mismatch',config=config(image_mode='ocr'))
