"""P-04 document adapters. Native Docling structures flow directly to the chunker.

Original binary input stays in the encrypted source revision. Locators refer
only to observed pages/regions, package XML, sheets/cells, or HWP sections.
Extraction and inference are not claimed to be verbatim original byte spans.
"""
from __future__ import annotations

import base64
import hashlib
import io
import json
import math
import re
import shutil
import subprocess
import zipfile
from collections import defaultdict
from functools import lru_cache
from importlib import metadata
from pathlib import Path, PurePosixPath

from airag import store
from airag.normalize import NormalizedDocument, OriginalAnchor

DOCX = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
HWP_TYPES = {'application/x-hwp', 'application/haansofthwp', 'application/vnd.hancom.hwp'}
HWPX_TYPES = {'application/hwp+zip', 'application/vnd.hancom.hwpx', 'application/x-hwpx'}
IMAGE_TYPES = {'image/png', 'image/jpeg', 'image/tiff', 'image/bmp'}
SUPPORTED = {'application/pdf', DOCX, XLSX} | HWP_TYPES | HWPX_TYPES | IMAGE_TYPES
DEFAULTS = {
    'adapter': 'documents_v1', 'adapter_version': 1,
    'pdf_artifacts_path': None, 'ocr_engine': 'tesseract', 'ocr_languages': ['ko', 'en'], 'tesseract_psm': 6, 'tessdata_path': None,
    'ocr_model_path': None, 'image_mode': 'ocr+vision', 'max_pages': 200,
    'max_expanded_bytes': 128 * 1024 * 1024, 'max_image_pixels': 20_000_000,
    'timeout_s': 120, 'hwp_classpath': None, 'hwp_sha256': None,
}


class DocumentError(store.ApiError):
    pass


def _file_digest(path):
    digest=hashlib.sha256()
    with Path(path).open('rb') as f:
        while block:=f.read(1024*1024):digest.update(block)
    return digest.hexdigest()


def canonical_config(config: dict) -> dict:
    if not isinstance(config, dict) or config.get('adapter') != 'documents_v1':
        raise DocumentError('DOCUMENT_PROFILE_REQUIRED', 'binary documents require an explicit documents_v1 normalizer profile', retryable=False)
    if set(config) - (set(DEFAULTS) | {'dependency_versions', 'model_fingerprints', 'vision_prompt_version'}):
        raise ValueError('unknown document adapter configuration field')
    cfg = {**DEFAULTS, **config}
    if type(cfg['adapter_version']) is not int or cfg['adapter_version'] != 1:
        raise ValueError('unsupported document adapter version')
    for name in ('max_pages', 'max_expanded_bytes', 'max_image_pixels', 'timeout_s'):
        value = cfg[name]
        if isinstance(value, bool) or not isinstance(value, int) or value <= 0:
            raise ValueError(f'{name} must be a positive integer')
    if cfg['ocr_engine'] not in ('tesseract', 'easyocr'):
        raise ValueError('ocr_engine must be tesseract or easyocr')
    if isinstance(cfg['tesseract_psm'], bool) or not isinstance(cfg['tesseract_psm'],int) or cfg['tesseract_psm'] not in (3,4,6,11,12,13):
        raise ValueError('tesseract_psm must select a supported text recognition mode')
    if cfg['image_mode'] not in ('ocr', 'vision', 'ocr+vision'):
        raise ValueError('image_mode must be ocr, vision, or ocr+vision')
    if not isinstance(cfg['ocr_languages'], list) or not cfg['ocr_languages'] or any(x not in ('ko', 'en') for x in cfg['ocr_languages']):
        raise ValueError('ocr_languages supports explicit ko/en language lists')
    for name in ('pdf_artifacts_path', 'ocr_model_path', 'hwp_classpath', 'hwp_sha256', 'tessdata_path'):
        if cfg[name] is not None and (not isinstance(cfg[name], str) or not cfg[name].strip()):
            raise ValueError(f'{name} must be a nonempty string or null')
    try:
        versions = {name: metadata.version(name) for name in ('docling-slim', 'docling-core', 'python-docx', 'openpyxl', 'Pillow')}
        if cfg['ocr_engine'] == 'easyocr':
            versions['easyocr'] = metadata.version('easyocr')
    except metadata.PackageNotFoundError as exc:
        raise DocumentError('DOCUMENT_DEPENDENCY_MISSING', 'install ai-feature-rag[documents] before using document adapters', retryable=False) from exc
    if cfg['ocr_engine']=='tesseract':
        versions['tesseract']='unavailable'
        if shutil.which('tesseract'):
            try:
                version=subprocess.run(['tesseract','--version'],capture_output=True,text=True,timeout=5,check=True)
                versions['tesseract']=version.stdout.splitlines()[0]
                args=['tesseract','--list-langs']
                if cfg['tessdata_path']:args+=['--tessdata-dir',cfg['tessdata_path']]
                langs=subprocess.run(args,capture_output=True,text=True,timeout=5,check=True)
                if cfg['tessdata_path'] is None:
                    match=re.search(r'List of available languages in "([^"]+)"',langs.stdout+langs.stderr)
                    if not match:raise ValueError('cannot resolve Tesseract data directory; configure tessdata_path explicitly')
                    cfg['tessdata_path']=match.group(1)
            except (OSError,subprocess.SubprocessError) as exc:
                raise DocumentError('OCR_RUNTIME_UNAVAILABLE','cannot inspect the configured Tesseract runtime',retryable=False) from exc
    if 'dependency_versions' in config and config['dependency_versions'] != versions:
        raise DocumentError('DOCUMENT_PROFILE_MISMATCH', 'installed document dependencies differ from the captured profile', retryable=False)
    cfg['dependency_versions'] = versions
    from airag.image_adapters import VISION_PROMPT_VERSION
    if config.get('vision_prompt_version', VISION_PROMPT_VERSION) != VISION_PROMPT_VERSION:
        raise DocumentError('DOCUMENT_PROFILE_MISMATCH','captured vision prompt version is not supported by this build',retryable=False)
    cfg['vision_prompt_version'] = VISION_PROMPT_VERSION
    fingerprints = {}
    if cfg['ocr_engine']=='tesseract' and cfg['tessdata_path']:
        names=[{'ko':'kor','en':'eng'}[x] for x in cfg['ocr_languages']]
        for name in names:
            path=Path(cfg['tessdata_path'])/(name+'.traineddata')
            if not path.is_file():raise DocumentError('OCR_MODELS_REQUIRED','configured Tesseract language data is missing',retryable=False)
            fingerprints[name+'.traineddata']=_file_digest(path)
        osd=Path(cfg['tessdata_path'])/'osd.traineddata'
        if osd.is_file():fingerprints['osd.traineddata']=_file_digest(osd)
    if cfg['pdf_artifacts_path']:
        for relative in ('docling-project--docling-layout-heron', 'docling-project--docling-models/model_artifacts/tableformer/accurate'):
            folder=Path(cfg['pdf_artifacts_path'])/relative
            files=sorted(p for p in folder.rglob('*') if p.is_file() and p.suffix in ('.json','.safetensors') and '.cache' not in p.parts)
            if not files:
                raise DocumentError('PDF_MODELS_REQUIRED','prepare the configured Docling layout/table models',retryable=False)
            digest=hashlib.sha256()
            for p in files:
                digest.update(str(p.relative_to(folder)).encode())
                with p.open('rb') as f:
                    while block:=f.read(1024*1024):digest.update(block)
            fingerprints[relative]=digest.hexdigest()
    if cfg['ocr_engine']=='easyocr' and cfg['ocr_model_path']:
        for name in ('craft_mlt_25k.pth','korean_g2.pth' if 'ko' in cfg['ocr_languages'] else 'english_g2.pth'):
            p=Path(cfg['ocr_model_path'])/name
            if not p.is_file():
                raise DocumentError('OCR_MODELS_REQUIRED','prepare the configured EasyOCR detection/recognition models',retryable=False)
            fingerprints[name]=_file_digest(p)
    if 'model_fingerprints' in config and config['model_fingerprints'] != fingerprints:
        raise DocumentError('DOCUMENT_PROFILE_MISMATCH','document model artifacts differ from the captured profile',retryable=False)
    cfg['model_fingerprints']=fingerprints
    if cfg['hwp_classpath']:
        jar = Path(cfg['hwp_classpath'])
        if not jar.is_file():
            raise ValueError('hwp_classpath must identify a local hwplib jar')
        actual = _file_digest(jar)
        if cfg['hwp_sha256'] is not None and cfg['hwp_sha256'] != actual:
            raise DocumentError('DOCUMENT_PROFILE_MISMATCH', 'HWP library digest differs from the captured profile', retryable=False)
        cfg['hwp_sha256'] = actual
    return cfg


def _zip(content: bytes, limit: int) -> zipfile.ZipFile:
    z = zipfile.ZipFile(io.BytesIO(content))
    infos = z.infolist()
    names = [x.filename for x in infos]
    if len(infos) > 10000 or len(names) != len(set(names)) or sum(x.file_size for x in infos) > limit:
        z.close()
        raise ValueError('archive exceeds document expansion bounds or has duplicate members')
    for info in infos:
        path = PurePosixPath(info.filename)
        if path.is_absolute() or '..' in path.parts or '\\' in info.filename or info.flag_bits & 1:
            z.close()
            raise ValueError('unsafe or encrypted document archive')
    return z


def _resources(z: zipfile.ZipFile, prefixes: tuple[str, ...]) -> list[dict]:
    resources = []
    for member in z.namelist():
        if member.startswith(prefixes) and not member.endswith('/'):
            data = z.read(member)
            resources.append({'member': member, 'sha256': hashlib.sha256(data).hexdigest(),
                              'byte_length': len(data), 'content_base64': base64.b64encode(data).decode()})
    return resources


def _item_text(item, doc) -> str:
    if hasattr(item, 'text'):
        return item.text
    if hasattr(item, 'export_to_markdown'):
        return item.export_to_markdown(doc=doc)
    return ''


def _anchor(text, locator, *, origin='extracted', granularity='document'):
    return OriginalAnchor(kind='document_location', original_text=text, locator=locator,
                          text_origin=origin, granularity=granularity)


@lru_cache(maxsize=4)
def _converter(config_json: str):
    from docling.document_converter import DocumentConverter, PdfFormatOption, WordFormatOption, ExcelFormatOption
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import PdfPipelineOptions, TesseractCliOcrOptions, EasyOcrOptions
    from docling.backend.msword_backend import MsWordDocumentBackend

    class AnchoredWordBackend(MsWordDocumentBackend):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self.docx_to_pdf_converter = None
            self.docx_to_pdf_converter_init = True

        def _handle_tables(self, element, doc):
            refs = super()._handle_tables(element, doc)
            if not hasattr(self, 'airag_tables'):
                self.airag_tables = []
            self.airag_tables.append((element, refs))
            return refs

    cfg = json.loads(config_json)
    options = {InputFormat.DOCX: WordFormatOption(backend=AnchoredWordBackend), InputFormat.XLSX: ExcelFormatOption()}
    if cfg['pdf_artifacts_path']:
        path = Path(cfg['pdf_artifacts_path'])
        if not path.is_dir():
            raise ValueError('pdf_artifacts_path must be an existing, explicitly prepared model directory')
        pdf = PdfPipelineOptions(artifacts_path=path, document_timeout=cfg['timeout_s'],
                                 do_ocr=True, do_table_structure=True, generate_picture_images=True, generate_parsed_pages=True,
                                 enable_remote_services=False, allow_external_plugins=False)
        if cfg['ocr_engine'] == 'tesseract':
            pdf.ocr_options = TesseractCliOcrOptions(lang=[{'ko':'kor','en':'eng'}[x] for x in cfg['ocr_languages']],psm=cfg['tesseract_psm'],path=cfg['tessdata_path'])
        else:
            if not cfg['ocr_model_path']:
                raise ValueError('EasyOCR requires an explicitly prepared ocr_model_path')
            pdf.ocr_options = EasyOcrOptions(lang=cfg['ocr_languages'], use_gpu=False,
                model_storage_directory=cfg['ocr_model_path'], download_enabled=False)
        options[InputFormat.PDF] = PdfFormatOption(pipeline_options=pdf)
    return DocumentConverter(allowed_formats=list(options), format_options=options)


def _office_pdf(media_type, content, name, cfg):
    from docling.datamodel.base_models import DocumentStream, ConversionStatus, InputFormat
    extension = {'application/pdf':'pdf', DOCX:'docx', XLSX:'xlsx'}[media_type]
    if media_type == 'application/pdf' and not cfg['pdf_artifacts_path']:
        raise DocumentError('PDF_MODELS_REQUIRED', 'prepare Docling layout/table models and configure pdf_artifacts_path', retryable=False)
    resources = []
    if media_type == 'application/pdf':
        import pypdfium2
        pdf = pypdfium2.PdfDocument(content)
        try:
            if len(pdf)>cfg['max_pages']:
                raise ValueError('PDF exceeds page bound')
            for i in range(len(pdf)):
                page=pdf[i]
                try:
                    width,height=page.get_size()
                    if width*height*9>cfg['max_image_pixels']:
                        raise ValueError('PDF page exceeds OCR rendering pixel bound')
                finally:
                    page.close()
        finally:
            pdf.close()
    if media_type in (DOCX, XLSX):
        with _zip(content, cfg['max_expanded_bytes']) as z:
            resources = _resources(z, ('word/media/', 'xl/media/'))
    if media_type == XLSX:
        import openpyxl
        workbook=openpyxl.load_workbook(io.BytesIO(content),read_only=True,keep_links=False)
        try:
            if len(workbook.worksheets)>cfg['max_pages'] or any((s.max_row or 0)*(s.max_column or 0)>1_000_000 for s in workbook.worksheets):
                raise ValueError('workbook exceeds sheet/cell bound')
        finally:
            workbook.close()
    result = _converter(json.dumps(cfg, sort_keys=True)).convert(
        DocumentStream(name=f'source.{extension}', stream=io.BytesIO(content)), raises_on_error=False,
        max_num_pages=cfg['max_pages'], max_file_size=store.MAX_CONTENT_BYTES,
    )
    if result.status != ConversionStatus.SUCCESS:
        raise DocumentError('DOCUMENT_CONVERSION_FAILED', f'{extension} conversion did not fully succeed ({result.status})', retryable=False)
    expected={'application/pdf':InputFormat.PDF,DOCX:InputFormat.DOCX,XLSX:InputFormat.XLSX}
    if result.input._backend.input_format!=expected[media_type]:
        raise DocumentError('MEDIA_TYPE_MISMATCH','declared media type differs from detected document format',retryable=False)
    doc = result.document
    doc.name = name
    if doc.origin is not None:
        doc.origin.filename = name
    anchors = {}
    artifacts = {'adapter':'documents_v1', 'media_type':media_type, 'dependency_versions':cfg['dependency_versions'], 'model_fingerprints':cfg['model_fingerprints'],
                 'resources':resources, 'warnings':[], 'pages':len(doc.pages)}
    page_sources={}
    for page in result.pages:
        parsed=page.parsed_page
        cells=(parsed.word_cells or parsed.textline_cells) if parsed is not None else []
        ocr=sum(bool(c.from_ocr) for c in cells)
        native=len(cells)-ocr
        page_sources[page.page_no]={'page_no':page.page_no,'native_text_cells':native,'ocr_text_cells':ocr}
    if media_type=='application/pdf':
        artifacts['page_text_sources']=list(page_sources.values())
    for item, _ in doc.iterate_items():
        if not hasattr(item, 'self_ref'):
            continue
        text = _item_text(item, doc)
        prov = [p.model_dump(mode='json') for p in getattr(item, 'prov', [])]
        locator = {'format':extension, 'docling_ref':item.self_ref, 'provenance':prov}
        origin = 'extracted'
        if media_type == 'application/pdf':
            sources=[page_sources.get(p['page_no'],{}) for p in prov]
            native=sum(p.get('native_text_cells',0) for p in sources)
            ocr=sum(p.get('ocr_text_cells',0) for p in sources)
            origin='ocr' if ocr and not native else 'extracted' if native and not ocr else 'extracted_or_ocr'
        anchors[item.self_ref] = _anchor(text, locator, origin=origin, granularity='page_region' if prov else 'document')
    if media_type == DOCX:
        _word_anchors(result.input._backend, doc, anchors, artifacts)
    elif media_type == XLSX:
        _excel_details(content, doc, anchors, artifacts, cfg)
    else:
        artifacts['ocr'] = {'engine':cfg['ocr_engine'], 'languages':cfg['ocr_languages'],
                            'tesseract_psm':cfg['tesseract_psm'] if cfg['ocr_engine']=='tesseract' else None,
                            'policy':'Docling DEFAULT OCR: detect bitmap regions and retain existing digital text',
                            'temporary_files':'Docling Tesseract CLI uses mode-0600 temporary PNG crops, removed in finally' if cfg['ocr_engine']=='tesseract' else 'EasyOCR receives page images in memory'}
        artifacts['warnings'].append('PDF extracted text and OCR share page-region provenance; no original byte-offset claim')
    return NormalizedDocument(doc, anchors, artifacts)


def _word_anchors(backend, doc, anchors, artifacts):
    parts = {}
    for part in backend.docx_obj.part.package.parts:
        root = getattr(part, 'element', None)
        if root is not None:
            parts[id(root)] = str(part.partname).lstrip('/')
    entries = list(backend.paragraph_to_items.items()) + getattr(backend, 'airag_tables', [])
    for element, refs in entries:
        root = element.getroottree().getroot()
        member = parts.get(id(root))
        if member is None:
            continue  # retain the coarser observed document locator
        locator = {'format':'docx', 'member':member, 'xpath':element.getroottree().getpath(element),
                   'namespaces':{k:v for k,v in element.nsmap.items() if k is not None}}
        for ref in refs:
            item = ref.resolve(doc)
            anchors[item.self_ref] = _anchor(_item_text(item, doc), locator, granularity='xml_element')
    artifacts['warnings'].append('DOCX locators are package XML elements, not rendered page coordinates')
    artifacts['warnings'].append('Automatic LibreOffice rendering is disabled; non-bitmap DrawingML shapes are retained in the original, not rasterized')


def _excel_details(content, doc, anchors, artifacts, cfg):
    import openpyxl
    from docling_core.types.doc import DocItemLabel
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=False, keep_links=False)
    cached = openpyxl.load_workbook(io.BytesIO(content), data_only=True, keep_links=False)
    sheets = []
    try:
        if len(wb.worksheets) > cfg['max_pages']:
            raise ValueError('workbook exceeds configured sheet bound')
        for index, sheet in enumerate(wb.worksheets, 1):
            if sheet.max_row * sheet.max_column > 1_000_000:
                raise ValueError('worksheet exceeds cell inspection bound')
            cells = []
            for row in sheet:
                for cell in row:
                    if cell.value is None:
                        continue
                    value = cell.value
                    if cell.data_type=='f' and not isinstance(value,str):
                        raise DocumentError('UNSUPPORTED_FORMULA','array/data-table formulas require a separate verified adapter; no recalculation fallback',retryable=False)
                    if hasattr(value, 'isoformat'):
                        value = value.isoformat()
                    if not isinstance(value, (str,int,float,bool)):
                        value = str(value)
                    if isinstance(value, float) and not math.isfinite(value):
                        raise ValueError('nonfinite spreadsheet value')
                    entry = {'cell':cell.coordinate, 'value':value, 'data_type':cell.data_type}
                    if cell.data_type == 'f':
                        entry['formula'] = value
                        cv = cached[sheet.title][cell.coordinate].value
                        entry['cached_value'] = cv if isinstance(cv,(str,int,float,bool,type(None))) else str(cv)
                        entry['recalculated'] = False
                        item = doc.add_text(DocItemLabel.TEXT, f'{sheet.title}!{cell.coordinate}: {value}')
                        anchors[item.self_ref] = _anchor(item.text, {'format':'xlsx','sheet':sheet.title,'cell':cell.coordinate}, granularity='cell')
                    cells.append(entry)
            sheets.append({'name':sheet.title,'sheet_index':index,'merged_ranges':[str(r) for r in sheet.merged_cells.ranges],'cells':cells})
        for item, _ in doc.iterate_items():
            if item.self_ref not in anchors:
                continue
            anchor = anchors[item.self_ref]
            if anchor.locator.get('cell'):
                continue
            prov = getattr(item, 'prov', [])
            if prov and 1 <= prov[0].page_no <= len(wb.worksheets):
                locator = {**anchor.locator, 'sheet':wb.worksheets[prov[0].page_no-1].title}
                anchors[item.self_ref] = _anchor(anchor.original_text, locator, granularity='sheet_region')
        artifacts['sheets'] = sheets
        artifacts['warnings'].append('Formula strings and cached values are preserved separately; no formula recalculation is performed')
    finally:
        wb.close(); cached.close()


def _xml(data):
    from defusedxml import ElementTree
    from defusedxml.common import DefusedXmlException
    from xml.etree.ElementTree import ParseError
    try:
        return ElementTree.fromstring(data)
    except (DefusedXmlException, ParseError) as exc:
        raise ValueError('unsafe or malformed document XML') from exc


def _tag(element):
    return element.tag.rsplit('}',1)[-1]


def _add_table(doc, cells):
    from docling_core.types.doc import TableData
    rows=max(c.end_row_offset_idx for c in cells)
    cols=max(c.end_col_offset_idx for c in cells)
    if rows<=0 or cols<=0 or rows*cols>1_000_000 or len(cells)>1_000_000:
        raise ValueError('table exceeds cell expansion bound')
    return doc.add_table(TableData(num_rows=rows,num_cols=cols,table_cells=cells))


def _hwpx(content, name, cfg):
    from docling_core.types.doc import DoclingDocument, DocItemLabel, TableData, TableCell
    doc = DoclingDocument(name=name)
    anchors = {}
    with _zip(content, cfg['max_expanded_bytes']) as z:
        if z.read('mimetype').strip()!=b'application/hwp+zip':
            raise ValueError('not an HWPX package')
        manifest=_xml(z.read('Contents/content.hpf'))
        entries=[e for e in manifest.iter() if _tag(e)=='item']
        members={e.get('id'):e.get('href') for e in entries}
        if len(entries)!=len(members) or None in members:
            raise ValueError('invalid HWPX manifest identifiers')
        section_members=[]
        for ref in (e for e in manifest.iter() if _tag(e)=='itemref'):
            member=members.get(ref.get('idref'))
            if not member or member not in z.namelist():
                raise ValueError('unresolved HWPX spine member')
            if member.startswith('Contents/section') and member.endswith('.xml'):
                section_members.append(member)
        actual={n for n in z.namelist() if n.startswith('Contents/section') and n.endswith('.xml')}
        if len(section_members)!=len(set(section_members)) or set(section_members)!=actual:
            raise ValueError('HWPX spine does not cover each section exactly once')
        if not section_members or len(section_members)>cfg['max_pages']:
            raise ValueError('HWPX must contain a bounded set of section XML members')
        resources = _resources(z, ('BinData/',))
        for si, member in enumerate(section_members):
            root = _xml(z.read(member))
            if root.tag!='{http://www.hancom.co.kr/hwpml/2011/section}sec':
                raise ValueError('invalid HWPX section root')
            def walk(element, pointer, in_table=False):
                tag = _tag(element)
                loc = {'format':'hwpx','member':member,'section':si,'element_path':pointer,'element_id':element.get('id')}
                if tag == 'tbl':
                    cells=[]
                    for ri, tr in enumerate(c for c in element if _tag(c)=='tr'):
                        ci=0
                        for tc in (c for c in tr if _tag(c)=='tc'):
                            addr=next((x for x in tc if _tag(x)=='cellAddr'),None)
                            span=next((x for x in tc if _tag(x)=='cellSpan'),None)
                            r=int(addr.get('rowAddr',ri)) if addr is not None else ri
                            c=int(addr.get('colAddr',ci)) if addr is not None else ci
                            rs=int(span.get('rowSpan',1)) if span is not None else 1
                            cs=int(span.get('colSpan',1)) if span is not None else 1
                            text='\n'.join(''.join(t.itertext()) for t in tc.iter() if _tag(t)=='t')
                            cells.append(TableCell(text=text,start_row_offset_idx=r,end_row_offset_idx=r+rs,start_col_offset_idx=c,end_col_offset_idx=c+cs,row_span=rs,col_span=cs))
                            ci=c+cs
                    if cells:
                        item=_add_table(doc,cells)
                        anchors[item.self_ref]=_anchor(_item_text(item,doc),loc,granularity='table')
                    return
                if tag=='p' and not in_table:
                    def texts(node):
                        if _tag(node)=='tbl': return
                        if node is not element and _tag(node)=='p': return
                        if _tag(node)=='t':
                            yield ''.join(node.itertext()); return
                        for child in node: yield from texts(child)
                    text=''.join(texts(element))
                    if text.strip():
                        item=doc.add_text(DocItemLabel.TEXT,text)
                        anchors[item.self_ref]=_anchor(text,loc,granularity='paragraph')
                counts=defaultdict(int)
                for child in element:
                    ct=_tag(child);counts[ct]+=1
                    walk(child,f'{pointer}/{ct}[{counts[ct]}]',in_table or tag=='tbl')
            walk(root,f'/{_tag(root)}[1]')
    return NormalizedDocument(doc,anchors,{'adapter':'hwpx_xml_v1','resources':resources,'sections':section_members,
        'warnings':['Section/paragraph/table XML locators are not rendered page coordinates']})


def _hwp(content, name, cfg):
    from docling_core.types.doc import DoclingDocument, DocItemLabel, TableData, TableCell
    if not cfg['hwp_classpath']:
        raise DocumentError('HWP_LIBRARY_REQUIRED','configure a pinned local hwplib jar in hwp_classpath',retryable=False)
    command=['java','-Xmx512m','--class-path',cfg['hwp_classpath'],str(Path(__file__).with_name('HwpBridge.java'))]
    proc=subprocess.run(command,input=content,capture_output=True,timeout=cfg['timeout_s'])
    if proc.returncode:
        raise DocumentError('HWP_CONVERSION_FAILED','HWP reader rejected input; encrypted/DRM/distribution or malformed files are unsupported',retryable=False)
    if len(proc.stdout)>cfg['max_expanded_bytes']:
        raise ValueError('HWP extraction exceeds expanded output bound')
    data=json.loads(proc.stdout)
    doc=DoclingDocument(name=name);anchors={}
    for block in data['blocks']:
        loc={'format':'hwp','path':block['path']}
        if block['kind']=='table':
            cells=[TableCell(text=c['text'],start_row_offset_idx=c['row'],end_row_offset_idx=c['row']+c['row_span'],start_col_offset_idx=c['col'],end_col_offset_idx=c['col']+c['col_span'],row_span=c['row_span'],col_span=c['col_span']) for c in block['cells']]
            if not cells: continue
            item=_add_table(doc,cells)
            anchors[item.self_ref]=_anchor(_item_text(item,doc),loc,granularity='table')
        elif block['text'].strip():
            item=doc.add_text(DocItemLabel.TEXT,block['text'])
            anchors[item.self_ref]=_anchor(item.text,loc,granularity='paragraph')
    return NormalizedDocument(doc,anchors,{'adapter':'hwplib','jar_sha256':cfg['hwp_sha256'],'resources':data['resources'],
        'warnings':['Section/paragraph/control locators only; hwplib does not provide rendered page numbers','Unsupported control types: '+', '.join(data['unsupported_controls'])]})


def build_document(media_type: str, content: bytes, *, name: str, config: dict, vision=None, heartbeat=None):
    if media_type not in SUPPORTED:
        raise ValueError(f'unsupported media_type for normalization: {media_type!r}')
    cfg=canonical_config(config)
    try:
        if media_type in ('application/pdf',DOCX,XLSX):
            return _office_pdf(media_type,content,name,cfg)
        if media_type in HWPX_TYPES:
            return _hwpx(content,name,cfg)
        if media_type in HWP_TYPES:
            return _hwp(content,name,cfg)
        from airag.image_adapters import build_image
        return build_image(content,name,cfg,vision,heartbeat,media_type)
    except store.ApiError:
        raise
    except (ValueError, OSError, RuntimeError, KeyError, zipfile.BadZipFile, subprocess.SubprocessError) as exc:
        raise DocumentError('DOCUMENT_CONVERSION_FAILED',f'{media_type} conversion failed ({type(exc).__name__})',retryable=False) from exc
